#!/usr/bin/env python3
"""Build the data-quality workbook from the import's output tables.

Usage: python3 build_report.py --out out --file "Engager import check.xlsx"
"""
import argparse
import csv
import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
INK = "17191C"
HEAD_FILL = PatternFill("solid", fgColor="17191C")
INPUT_FILL = PatternFill("solid", fgColor="FFF2A8")   # your decisions
AMBER_FILL = PatternFill("solid", fgColor="FBEEDC")
RED_FILL = PatternFill("solid", fgColor="FBE5E2")
GREEN = "2E6A3E"
THIN = Side(style="thin", color="DFDDD6")


def read(out, name):
    with open(out / f"{name}.csv", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def d(s):
    return dt.date.fromisoformat(s) if s else None


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def body(cell, bold=False, color=INK, wrap=True):
    cell.font = Font(name=FONT, size=10, bold=bold, color=color)
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    cell.border = Border(bottom=THIN)


def table(ws, headers, rows, widths, start=1, inputs=()):
    for i, h in enumerate(headers, 1):
        ws.cell(row=start, column=i, value=h)
    style_header(ws, start, len(headers))
    for r, row in enumerate(rows, start + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            body(cell, bold=(c == 1))
            if c in inputs:
                cell.fill = INPUT_FILL
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    if rows:
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"
    return start + len(rows)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--file", required=True)
    a = ap.parse_args()
    out = Path(a.out)

    summary = json.loads((out / "summary.json").read_text())
    clients = read(out, "clients")
    jobs = read(out, "jobs")
    issues = read(out, "quality_issues")
    groups = {g["id"]: g for g in read(out, "groups_suggested")}
    members = read(out, "group_members_suggested")
    services = read(out, "client_services")
    names = {c["id"]: c["name"] for c in clients}
    unchecked = [c["name"] for c in clients if c["kind"] in ("Ltd", "LLP") and c.get("ch_checked") != "True"]

    wb = Workbook()

    # ---------------- Summary ----------------
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 70
    ws["A1"] = "Engager import check"
    ws["A1"].font = Font(name=FONT, size=18, bold=True, color=INK)
    ws["A2"] = "What came across from Engager, and what needs your decision before go-live."
    ws["A2"].font = Font(name=FONT, size=10, color="57606A")
    ws["A4"] = "As at"
    ws["B4"] = d(summary["today"])
    ws["B4"].number_format = "d mmm yyyy"
    ws["C4"] = "Days-left figures on the Open jobs sheet count from this date."
    for c in ("A4", "B4", "C4"):
        body(ws[c], bold=(c == "A4"))

    rows = [
        ("Clients imported", summary["clients"],
         f"{summary['by_kind']['Ltd']} companies, {summary['by_kind']['LLP']} LLP, "
         f"{summary['by_kind']['Individual']} individuals"),
        ("Contacts (not clients)", summary["contacts"], "People linked to clients, imported as contacts"),
        ("Open jobs", summary["open_jobs"], "From the Engager jobs list"),
        ("Completed jobs (history)", summary["history_jobs"],
         f"Back to 2023. {len(summary['history_unmatched_clients'])} former clients kept as names only."),
        ("Groups suggested", summary["groups_suggested"], "From Engager's relationships; confirm on the Groups sheet"),
        ("Needs action now", "='Action now'!$F$1", "Late or close to deadline, or missing at Companies House"),
        ("To check", "='Checks'!$F$1", "Jobs to tick off, missing jobs, deadline mismatches"),
        ("Clients missing a reference or contact", "='Client records'!$N$1",
         "Company number, auth code, UTRs, NI, VAT or PAYE refs, email or mobile"),
        ("No signed engagement letter",
         "=COUNTIFS('Client records'!$K$2:$K$500,\"<>Signed\",'Client records'!$A$2:$A$500,\"<>\")",
         "Engager shows only 3 signed; the rest are none on file or sent but not signed"),
        ("AML not recorded in Xama", "=COUNTIF('Client records'!$L$2:$L$500,\"Not linked\")"
                                     "+COUNTIF('Client records'!$L$2:$L$500,\"No assessment\")",
         "Not linked to Xama, or no risk assessment yet"),
    ]
    ws["A6"], ws["B6"], ws["C6"] = "Item", "Count", "Notes"
    style_header(ws, 6, 3)
    for i, (label, value, note) in enumerate(rows, 7):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
        ws.cell(row=i, column=3, value=note)
        for c in range(1, 4):
            body(ws.cell(row=i, column=c), bold=(c == 1))
        ws.cell(row=i, column=2).alignment = Alignment(horizontal="right", vertical="top")

    r = 7 + len(rows) + 1
    notes = [
        ("Not imported", "Engager's business profile fields (VAT quarterly, payroll monthly, Nest, CIS nil, "
                         "QBO Essentials, Xero, 3-month turnaround, Dext). They are identical on all 52 records, "
                         "so they are form defaults, not real information. Each client's services come from "
                         "their service set-up instead. Also left out: 2022/23 income figures and the AML "
                         "questionnaire answers (these stay in Xama)."),
        ("Time budgets", f"Only {summary['jobs_with_budget']} open jobs have a time budget, and no completed job "
                         "has time recorded. Estimates per job type need setting before capacity planning "
                         "is accurate."),
        ("Suggested slots", f"{summary['jobs_with_suggested_slot']} open jobs have a suggested planning date, "
                            "based on when you did the same job last time (at least a week before the deadline)."),
        ("Companies House", "Each company's deadlines were checked against Companies House on "
                            f"{d(summary['today']):%-d %b %Y}. "
                            + (f"Not checked: {', '.join(unchecked)}; the daily check will cover them."
                               if unchecked else "Every company was checked.")),
        ("Name fixes", "Company names with title-cased initials have been corrected from the name-fixes list. "
                       "Titles such as Mr are kept separately."),
    ]
    ws.cell(row=r, column=1, value="Notes")
    ws.cell(row=r, column=1).font = Font(name=FONT, size=12, bold=True, color=INK)
    for i, (label, text) in enumerate(notes, r + 1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=3, value=text)
        body(ws.cell(row=i, column=1), bold=True)
        body(ws.cell(row=i, column=3))
        ws.row_dimensions[i].height = 15 * max(2, len(text) // 75 + 1)

    r = r + len(notes) + 2
    ws.cell(row=r, column=1, value="How to use this workbook")
    ws.cell(row=r, column=1).font = Font(name=FONT, size=12, bold=True, color=INK)
    guide = [
        ("Yellow cells", "are yours to fill in. Pick from the drop-down where there is one."),
        ("Example", "Action now › [client] › Your decision: 'Strike-off intended' · "
                    "Notes: 'Client closing the company; DS01 filed 12 Sep'."),
        ("Then", "send the workbook back. Your decisions are applied when the data goes into SharePoint."),
    ]
    for i, (label, text) in enumerate(guide, r + 1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=3, value=text)
        body(ws.cell(row=i, column=1), bold=True)
        body(ws.cell(row=i, column=3))
        if label == "Yellow cells":
            ws.cell(row=i, column=2).fill = INPUT_FILL

    # ---------------- Action now ----------------
    ws = wb.create_sheet("Action now")
    act = [x for x in issues if x["level"] == "Action now"]
    end = table(ws, ["Client", "What", "Detail", "Your decision", "Notes"],
                [[x["client"], x["issue"], x["detail"], None, None] for x in act],
                [34, 52, 64, 24, 36], inputs=(4, 5))
    dv = DataValidation(type="list", allow_blank=True,
                        formula1='"Done - mark complete,In hand,Needs doing,Strike-off intended,Not our job"')
    ws.add_data_validation(dv)
    if act:
        dv.add(f"D2:D{end}")
    ws["F1"] = f"=COUNTA(A2:A{max(end, 2)})"
    ws["F1"].font = Font(name=FONT, size=10, color="FFFFFF")
    ws.column_dimensions["F"].hidden = True

    # ---------------- Checks ----------------
    ws = wb.create_sheet("Checks")
    chk = [x for x in issues if x["level"] == "Check" and x["area"] != "Compliance"]
    area_order = {"Companies House": 0, "Deadlines": 1, "Jobs": 2}
    chk.sort(key=lambda x: (area_order.get(x["area"], 9), x["client"]))
    end = table(ws, ["Client", "Area", "What", "Detail", "Your decision"],
                [[x["client"], x["area"], x["issue"], x["detail"], None] for x in chk],
                [34, 16, 56, 64, 24], inputs=(5,))
    dv = DataValidation(type="list", allow_blank=True,
                        formula1='"Done - mark complete,Set up the job,Use Companies House date,Not needed,Leave as is"')
    ws.add_data_validation(dv)
    if chk:
        dv.add(f"E2:E{end}")
    ws["F1"] = f"=COUNTA(A2:A{max(end, 2)})"
    ws["F1"].font = Font(name=FONT, size=10, color="FFFFFF")
    ws.column_dimensions["F"].hidden = True

    # ---------------- Client records ----------------
    ws = wb.create_sheet("Client records")
    svc_by = {}
    for s in services:
        svc_by.setdefault(s["client_id"], set()).add(s["service"])
    for j in jobs:
        if j["client_id"]:
            svc_by.setdefault(j["client_id"], set()).add(j["service"])

    headers = ["Client", "Type", "Company no.", "CH auth code", "CT UTR", "SA UTR", "NI number",
               "VAT number", "PAYE ref", "Email or mobile", "Engagement letter", "AML (Xama)", "Gaps"]
    rows_out, flags = [], []
    for c in sorted(clients, key=lambda c: (c["kind"] == "Individual", c["name"])):
        svc = svc_by.get(c["id"], set())
        company = c["kind"] in ("Ltd", "LLP")
        sa = "SA100" in svc

        def need(present, required):
            return ("✓" if present else "Missing") if required else ("✓" if present else "–")

        loe = {"Finalised": "Signed", "Active": "Sent, not signed"}.get(c["loe_status"], "None on file")
        aml = {"not linked": "Not linked", "no assessment": "No assessment",
               "in progress": "In progress"}.get(c["xama_status"] or "", c["xama_status"] or "Not linked")
        row = [c["name"], c["kind"],
               need(c["company_number"], company),
               need(c["ch_auth_code"], c["kind"] == "Ltd"),
               need(c["ct_utr"], c["kind"] == "Ltd"),
               need(c["utr"], sa),
               need(c["ni_number"], sa),
               need(c["vat_number"], "VAT" in svc),
               need(c["paye_ref"], "PAYROLL" in svc),
               need(c["email"] or c["mobile"], True),
               loe, aml, None]
        rows_out.append(row)
    end = table(ws, headers, rows_out, [36, 11, 12, 12, 10, 10, 11, 11, 10, 13, 17, 15, 8])
    for r in range(2, end + 1):
        for c in range(3, 13):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            v = cell.value
            if v == "Missing" or v in ("None on file", "Not linked", "No assessment"):
                cell.fill = AMBER_FILL
                cell.font = Font(name=FONT, size=10, bold=True, color="7A4206")
            elif v == "✓" or v == "Signed":
                cell.font = Font(name=FONT, size=10, color=GREEN)
        ws.cell(row=r, column=13, value=(
            f'=COUNTIF(C{r}:J{r},"Missing")+IF(K{r}="Signed",0,1)'
            f'+IF(OR(L{r}="Not linked",L{r}="No assessment"),1,0)'))
        body(ws.cell(row=r, column=13))
        ws.cell(row=r, column=13).alignment = Alignment(horizontal="center", vertical="top")
    for r in range(2, end + 1):
        ws.cell(row=r, column=15, value=f'=IF(COUNTIF(C{r}:J{r},"Missing")>0,1,0)')
    ws["N1"] = f"=SUM(O2:O{end})"
    ws["N1"].font = Font(name=FONT, size=10, color="FFFFFF")
    ws.column_dimensions["N"].hidden = True
    ws.column_dimensions["O"].hidden = True
    ws.cell(row=1, column=13).comment = Comment(
        "Count of missing references, plus 1 if no signed engagement letter, plus 1 if AML not recorded.",
        "Import")

    # ---------------- Groups ----------------
    ws = wb.create_sheet("Groups")
    rows_out = []
    for m in members:
        g = groups[m["group_id"]]
        rows_out.append([g["name"], m["client_name"],
                         {"Ltd": "Company", "LLP": "LLP"}.get(m["kind"], m["kind"]), None, None])
    end = table(ws, ["Suggested group", "Member", "Type", "Keep in group? (Y/N)", "Rename group to"],
                rows_out, [30, 38, 12, 20, 28], inputs=(4, 5))
    dv = DataValidation(type="list", allow_blank=True, formula1='"Y,N"')
    ws.add_data_validation(dv)
    dv.add(f"D2:D{end}")
    ws.cell(row=end + 2, column=1,
            value="Built from Engager's Relationships, using primary contacts and company links only. "
                  "Anyone listed only as a secondary contact was left out of the group; "
                  "add them by typing a new row.")
    body(ws.cell(row=end + 2, column=1))
    ws.merge_cells(start_row=end + 2, start_column=1, end_row=end + 2, end_column=5)
    ws.row_dimensions[end + 2].height = 42

    # ---------------- Open jobs ----------------
    ws = wb.create_sheet("Open jobs")
    jrows = sorted(jobs, key=lambda j: (j["deadline"] or "9999", j["client_name"]))
    ch_label = {"filed": "Already filed", "matches": "Matches", "differs": "Differs"}
    rows_out = []
    for j in jrows:
        rows_out.append([j["client_name"], j["title"], d(j["period_end"]), j["stage_name"], d(j["deadline"]),
                         None, ch_label.get(j["ch_check"], ""), d(j["ch_deadline"]),
                         d(j["suggested_slot"]), d(j["last_completed"]),
                         float(j["estimate_hours"]) if j["estimate_hours"] not in ("", "0.0") else None])
    end = table(ws, ["Client", "Job", "Period end", "Current stage", "Deadline", "Days left",
                     "Companies House", "CH deadline", "Suggested slot", "Last done", "Budget (hours)"],
                rows_out, [34, 30, 12, 28, 12, 10, 15, 12, 13, 12, 11])
    for r in range(2, end + 1):
        for c in (3, 5, 8, 9, 10):
            ws.cell(row=r, column=c).number_format = "d mmm yyyy"
        ws.cell(row=r, column=6, value=f'=IF(E{r}="","",E{r}-Summary!$B$4)')
        body(ws.cell(row=r, column=6))
        ws.cell(row=r, column=6).number_format = '0;[Red]-0'
        ws.cell(row=r, column=11).number_format = "0.00"
        v = ws.cell(row=r, column=7).value
        if v == "Already filed":
            ws.cell(row=r, column=7).font = Font(name=FONT, size=10, bold=True, color=GREEN)
        elif v == "Differs":
            ws.cell(row=r, column=7).fill = AMBER_FILL

    wb.save(a.file)
    print("saved", a.file)


if __name__ == "__main__":
    main()
